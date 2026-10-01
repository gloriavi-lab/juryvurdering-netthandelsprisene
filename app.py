"""
Juryvurdering Netthandelsprisene – Ekspertvurdering (Fase 2)
==============================================================
Rendyrket, lett vurderingsverktøy for ekspertvurderingen (Fase 2) i
Netthandelsprisene. Bygget som en EGEN app – bevisst adskilt fra
hovedappen (screening + database over alle nettbutikker) – fordi
hovedappen ble treg av å håndtere 1000 butikker og mange faner i ett
og samme Google Sheet.

Arbeidsflyt:
1. Fase 1 gjøres i et eget Excel-ark (utenfor denne appen). Topp ~50
   derfra lastes opp her (sidepanelet) som butikklisten for fasen.
2. Hvert jurymedlem skriver navnet sitt og velger sitt fagfelt (samme
   5 kategorier som i Fase 1).
3. Jurymedlemmet vurderer butikkene på KUN sitt fagfelt, med akkurat
   de samme kriteriene (ord for ord) som ble brukt i Fase 1 – pluss et
   kommentarfelt per kriterium, som ikke fantes i Fase 1-arket.
4. Lagring skjer i et eget, lite Google Sheet (IKKE hovedappens store
   ark). Appen skriver KUN raden(e) for butikken som akkurat ble
   lagret – aldri hele arket – slik at manuelle endringer gjort
   direkte i Google Sheets alltid overlever og vises i appen.
5. "Finale"-fanen rangerer butikkene (snitt av alle registrerte
   ekspertscorer) gruppert på størrelsesklasse, og viser alle
   kommentarer som er lagt inn – til bruk når finalistene skal pekes ut.
"""

import re
from collections import defaultdict
from datetime import datetime

import gspread
import openpyxl
import streamlit as st
from google.oauth2.service_account import Credentials

st.set_page_config(page_title="Ekspertvurdering – Netthandelsprisene", page_icon="🎓", layout="wide")

st.markdown("""
<style>
.stApp { background-color: #FFFFFF; }
section[data-testid="stSidebar"] > div { background-color: #FFFFFF !important; border-right: 1px solid #E8E6E2; }
div[data-testid="stExpander"] { background-color: white; border-radius: 8px; border: 1px solid #E8E6E2; }
.main-header { background: #212121; color: white; padding: 16px 28px; border-radius: 10px; margin-bottom: 20px; display: flex; align-items: center; gap: 16px; }
.logo-badge { background: #C8102E; color: white; padding: 5px 12px; border-radius: 4px; font-weight: 700; font-size: 13px; text-transform: uppercase; letter-spacing: 0.5px; }
.klasse-card { background: #F8F7F5; border-radius: 10px; padding: 18px 20px; border-top: 4px solid #C8102E; height: 100%; }
.klasse-card h4 { margin: 0 0 6px 0; }
.klasse-card .belop { color: #555; font-size: 14px; }
.kat-pill { display: inline-block; background: #C8102E; color: white; font-weight: 700; padding: 6px 16px; border-radius: 20px; font-size: 14px; margin-bottom: 10px; }
.kat-card { background: #F8F7F5; border-radius: 10px; padding: 16px 18px; margin-bottom: 14px; border: 1px solid #E8E6E2; }
.kat-card ul { margin: 6px 0 0 0; padding-left: 20px; }
.kat-card li { margin-bottom: 4px; color: #333; }
.finale-rad { background: #F8F7F5; border-radius: 10px; padding: 14px 18px; margin-bottom: 10px; border-left: 4px solid #C8102E; }
.finale-plass { font-weight: 800; font-size: 20px; color: #C8102E; margin-right: 10px; }
.snitt-badge { background: #212121; color: white; padding: 3px 10px; border-radius: 6px; font-size: 13px; font-weight: 600; }
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
# Kriterier – IDENTISK ordlyd som i Fase 1-arket (5 kategorier, 16 kriterier)
# ─────────────────────────────────────────────
KRITERIER = [
    ("Førsteinntrykk", "Opplevelse: Bildebruk, beskrivelser o.l."),
    ("Førsteinntrykk", "Navigasjon / UX"),
    ("Førsteinntrykk", "Søk: auto-korrektur, synonymer, utlisting, forslag"),
    ("Kundeservice & Tilgjengelighet", "Lett tilgjengelig info om levering, retur og kjøpsvilkår"),
    ("Kundeservice & Tilgjengelighet", "Google/Trustpilot/andre åpne løsninger for kundetilfredshet"),
    ("Kundeservice & Tilgjengelighet", "Tilgjengelighet, åpningstider, kanaler, responstid"),
    ("Kundeservice & Tilgjengelighet", "Bærekraft"),
    ("Kjøp/inspirasjon/personalisering", "Kassen: Levering, betaling"),
    ("Kjøp/inspirasjon/personalisering", "Inspirasjon"),
    ("Kjøp/inspirasjon/personalisering", "Mersalg/anbefalinger/personalisering"),
    ("Markedsføring/kundedialog", "Bruk av SoMe"),
    ("Markedsføring/kundedialog", "Betalt og organisk synlighet"),
    ("Markedsføring/kundedialog", "E-post"),
    ("Innovasjon", "Innovative løsninger"),
    ("Innovasjon", "Adopsjon av nye teknologier for styrke konkurransekraft"),
    ("Innovasjon", "Kommersielt håndverk"),
]
FAGFELT = list(dict.fromkeys(kat for kat, _ in KRITERIER))
KRITERIER_PER_FAGFELT = {kat: [krit for k, krit in KRITERIER if k == kat] for kat in FAGFELT}

# Størrelsesklasser – definisjon avtalt med oppdragsgiver
KLASSE_DEFINISJON = [
    ("Liten", "Under 50 mill kr"),
    ("Medium", "50–250 mill kr"),
    ("Stor", "Over 250 mill kr"),
]
KLASSER = [k for k, _ in KLASSE_DEFINISJON]

# Forslag til fagfelt per jurymedlem (brukes kun til å forhåndsvelge i dropdown –
# jurymedlemmet kan alltid overstyre selv). Matcher på om navnet DE skriver inn
# inneholder nøkkelordet (små bokstaver), så "Ole Johan H." treffer "ole johan".
JURY_FAGFELT_FORSLAG = {
    "ole johan": "Førsteinntrykk",
    "stian": "Kundeservice & Tilgjengelighet",
    "torkel": "Kjøp/inspirasjon/personalisering",
    "marte": "Kjøp/inspirasjon/personalisering",
    "vikki": "Markedsføring/kundedialog",
    "nicholas": "Markedsføring/kundedialog",
    "guro": "Innovasjon",
}

VURDERINGER_HEADER = ["Butikk", "Jurymedlem", "Fagfelt", "Kriterium", "Score", "Kommentar", "Sist oppdatert"]
BUTIKKLISTE_FIL = "butikkliste.json"


# ─────────────────────────────────────────────
# Lokal lagring av opplastet butikkliste (overlever refresh, ikke deploy-restart)
# ─────────────────────────────────────────────
def lagre_lokalt(data):
    import json
    with open(BUTIKKLISTE_FIL, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)


def last_lokalt():
    import json, os
    if not os.path.exists(BUTIKKLISTE_FIL):
        return None
    try:
        with open(BUTIKKLISTE_FIL, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


# ─────────────────────────────────────────────
# Google Sheets – eget, lite regneark
# ─────────────────────────────────────────────
@st.cache_resource(ttl=3600, show_spinner=False)
def _gc_klient():
    try:
        creds_info = st.secrets["gcp_service_account"]
        scope = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
        creds = Credentials.from_service_account_info(creds_info, scopes=scope)
        return gspread.authorize(creds)
    except Exception as e:
        st.session_state["_gsheets_feil"] = str(e)
        return None


@st.cache_resource(ttl=3600, show_spinner=False)
def koble_sheets():
    gc = _gc_klient()
    if not gc:
        return None
    try:
        return gc.open_by_key(st.secrets["google_sheets"]["sheet_id"])
    except Exception as e:
        st.session_state["_gsheets_feil"] = str(e)
        return None


def hent_vurderinger_ark(sh):
    """Henter "Vurderinger"-fanen, oppretter den KUN hvis den ikke finnes fra før –
    aldri gjenoppbygger en fane som allerede eksisterer (det var det som gjorde at
    manuelle Sheets-endringer forsvant i den gamle appen)."""
    try:
        return sh.worksheet("Vurderinger")
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title="Vurderinger", rows=1000, cols=len(VURDERINGER_HEADER))
        ws.append_row(VURDERINGER_HEADER)
        return ws


@st.cache_data(ttl=15, show_spinner=False)
def hent_vurderinger(_sh):
    """Leser rå verdier direkte fra arket (kort mellomlagring på 15 sek, kun for å
    unngå å sprenge API-kvoten ved rask klikking – manuelle Sheets-endringer dukker
    uansett opp i appen i løpet av maks 15 sekunder)."""
    ws = hent_vurderinger_ark(_sh)
    return ws.get_all_values()


def lagre_vurderinger_batch(sh, butikk, jurymedlem, fagfelt, vurderinger):
    """vurderinger: liste av (kriterium, score, kommentar).
    Skriver KUN radene som gjelder akkurat denne butikken + dette jurymedlemmet –
    ALDRI hele arket – slik at manuelle endringer andre gjør direkte i Google Sheets
    ikke blir overskrevet av appen."""
    ws = hent_vurderinger_ark(sh)
    rå = ws.get_all_values()
    rader = rå[1:] if rå else []
    nokkel_til_rad = {}
    for i, rad in enumerate(rader, start=2):
        if len(rad) >= 4:
            nokkel_til_rad[(rad[0], rad[1], rad[3])] = i

    tidsstempel = datetime.now().strftime("%Y-%m-%d %H:%M")
    oppdateringer = []
    nye_rader = []
    for kriterium, score, kommentar in vurderinger:
        nokkel = (butikk, jurymedlem, kriterium)
        rad_data = [butikk, jurymedlem, fagfelt, kriterium, str(score), kommentar, tidsstempel]
        if nokkel in nokkel_til_rad:
            radnr = nokkel_til_rad[nokkel]
            oppdateringer.append({"range": f"A{radnr}:G{radnr}", "values": [rad_data]})
        else:
            nye_rader.append(rad_data)

    if oppdateringer:
        ws.batch_update(oppdateringer)
    if nye_rader:
        ws.append_rows(nye_rader)
    hent_vurderinger.clear()


# ─────────────────────────────────────────────
# Import av Fase 1-Excel-fila → butikkliste for denne fasen
# ─────────────────────────────────────────────
def les_fase1_liste(opplastet_fil):
    """Leser butikkliste fra Fase 1-arket. Forventer en rad med "Butikk" i kolonne A
    (samme header-rad som Fase 1-malen "Netthandelsprisene_Fase 1.xlsx"), og kolonner
    "Klasse", "Bransje", "URL" et sted i samme rad (rekkefølge spiller ingen rolle).
    URL-kolonnen kan enten være ren tekst eller en =HYPERLINK(...)-formel."""
    try:
        wb = openpyxl.load_workbook(opplastet_fil, data_only=False)
    except Exception as e:
        st.error(f"Kunne ikke lese Excel-filen: {e}")
        return None

    ws = wb.worksheets[0]
    header_rad = None
    for r in range(1, min(ws.max_row, 10) + 1):
        if str(ws.cell(row=r, column=1).value or "").strip().lower() == "butikk":
            header_rad = r
            break
    if header_rad is None:
        st.error('Fant ikke en rad med "Butikk" i kolonne A i fila.')
        return None

    kolonner = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_rad, column=c).value
        if v:
            kolonner[str(v).strip().lower()] = c

    kol_butikk = kolonner.get("butikk")
    kol_klasse = kolonner.get("klasse")
    kol_bransje = kolonner.get("bransje")
    kol_url = kolonner.get("url")

    butikker = {}
    for r in range(header_rad + 1, ws.max_row + 1):
        navn_celle = ws.cell(row=r, column=kol_butikk).value
        if not navn_celle or not str(navn_celle).strip():
            continue
        navn = str(navn_celle).strip()

        url = ""
        if kol_url:
            url_verdi = ws.cell(row=r, column=kol_url).value
            if url_verdi:
                treff = re.search(r'HYPERLINK\("([^"]+)"', str(url_verdi))
                url = treff.group(1) if treff else str(url_verdi).strip()

        butikker[navn] = {
            "klasse": str(ws.cell(row=r, column=kol_klasse).value or "").strip() if kol_klasse else "",
            "bransje": str(ws.cell(row=r, column=kol_bransje).value or "").strip() if kol_bransje else "",
            "url": url,
        }
    return butikker


# ─────────────────────────────────────────────
# Finale – aggregering av alle registrerte ekspertscorer
# ─────────────────────────────────────────────
def beregn_finale_data(alle_rader, butikkliste):
    """Returnerer:
    - snitt_per_butikk: {butikk: (snittscore, antall_scorer)}
    - detaljer_per_butikk: {butikk: [(fagfelt, kriterium, score, kommentar, jurymedlem), ...]}
    Bruker ALLE registrerte vurderinger uansett jurymedlem – finalen skal reflektere
    summen av ekspertenes vurderinger, ikke kun én person sine."""
    detaljer = defaultdict(list)
    scorer = defaultdict(list)
    if len(alle_rader) > 1:
        for rad in alle_rader[1:]:
            if len(rad) < 6:
                continue
            butikk, jurymedlem, fagfelt, kriterium, score, kommentar = rad[:6]
            try:
                score_num = float(score)
            except (ValueError, TypeError):
                continue
            scorer[butikk].append(score_num)
            detaljer[butikk].append((fagfelt, kriterium, score_num, kommentar, jurymedlem))
    snitt = {b: (sum(v) / len(v), len(v)) for b, v in scorer.items()}
    return snitt, detaljer


# ─────────────────────────────────────────────
# UI – sidepanel (felles for alle faner)
# ─────────────────────────────────────────────
st.markdown(
    '<div class="main-header"><span class="logo-badge">Netthandelsprisene</span>'
    '<span style="font-size:20px;font-weight:700;">🎓 Ekspertvurdering – Fase 2</span></div>',
    unsafe_allow_html=True,
)

sh = koble_sheets()

with st.sidebar:
    side = st.radio("Naviger", ["📖 Oversikt", "✍️ Vurdering", "🏆 Finale"], label_visibility="collapsed")
    st.markdown("---")

    st.subheader("⚙️ Oppsett")
    if sh:
        st.success("✅ Koblet til Google Sheets")
        st.markdown(f"[📊 Åpne arket]({sh.url})")
    else:
        st.error("❌ Ikke koblet til Google Sheets ennå")
        with st.expander("Feilmelding"):
            st.code(st.session_state.get("_gsheets_feil", "Ukjent feil"))

    st.markdown("---")
    st.subheader("📤 Butikkliste for denne fasen")
    if "butikkliste" not in st.session_state:
        st.session_state.butikkliste = last_lokalt()

    if st.session_state.butikkliste:
        st.success(f"{len(st.session_state.butikkliste)} butikker lastet inn.")
        if st.button("🗑️ Fjern liste"):
            st.session_state.butikkliste = None
            lagre_lokalt(None)
            st.rerun()

    opplastet = st.file_uploader("Last opp Fase 1-Excel (.xlsx)", type=["xlsx"])
    if opplastet is not None:
        ny_liste = les_fase1_liste(opplastet)
        if ny_liste:
            st.session_state.butikkliste = ny_liste
            lagre_lokalt(ny_liste)
            st.success(f"✅ {len(ny_liste)} butikker lastet inn!")
            st.rerun()

butikkliste = st.session_state.get("butikkliste")


# ─────────────────────────────────────────────
# Side: Oversikt
# ─────────────────────────────────────────────
def vis_oversikt():
    st.header("📖 Kriterier og retningslinjer")
    st.caption("Samme kriterier som i Fase 1 – hvert kriterium vurderes fra 1 til 5 stjerner (1 = svak, 5 = utmerket), med mulighet for å legge inn en kommentar.")

    st.subheader("Størrelsesklasser")
    kol = st.columns(3)
    for c, (klasse, belop) in zip(kol, KLASSE_DEFINISJON):
        c.markdown(f'<div class="klasse-card"><h4>{klasse}</h4><div class="belop">{belop}</div></div>', unsafe_allow_html=True)

    st.markdown("")
    st.subheader("Kriterier per fagfelt")
    kol2 = st.columns(len(FAGFELT))
    for c, kat in zip(kol2, FAGFELT):
        punkter = "".join(f"<li>{krit}</li>" for krit in KRITERIER_PER_FAGFELT[kat])
        c.markdown(
            f'<div class="kat-card"><span class="kat-pill">{kat}</span><ul>{punkter}</ul></div>',
            unsafe_allow_html=True,
        )

    st.markdown("")
    st.subheader("Veien til finalen")
    st.markdown(
        "- **Fase 1:** Alle nominerte butikker vurderes av juryen. Topp ca. 50 (etter snittscore) går videre.\n"
        "- **Fase 2 – Ekspertvurdering:** Hvert jurymedlem vurderer butikkene på sitt eget fagfelt.\n"
        "- **Finale:** Butikker som har vært gjennom *både* Fase 1 og Fase 2 rangeres etter snittscore, "
        "og de beste per størrelsesklasse (Liten / Medium / Stor) blir finalister."
    )


# ─────────────────────────────────────────────
# Side: Vurdering
# ─────────────────────────────────────────────
def vis_vurdering(sh, butikkliste):
    st.header("✍️ Vurdering")

    if not sh:
        st.warning("Google Sheets-tilkoblingen mangler – sjekk `secrets` i sidepanelet før vurderinger kan lagres.")
        return
    if not butikkliste:
        st.info("💡 Last opp Fase 1-Excel-fila i sidepanelet til venstre for å komme i gang.")
        return

    col1, col2 = st.columns(2)
    with col1:
        jurynavn = st.text_input("Ditt navn (jurymedlem)", value=st.session_state.get("_jurynavn", ""))
        st.session_state["_jurynavn"] = jurynavn
    with col2:
        forslag = next((v for k, v in JURY_FAGFELT_FORSLAG.items() if k in jurynavn.lower()), FAGFELT[0])
        fagfelt_valgt = st.selectbox(
            "Ditt fagfelt (kategorien du er ekspert på)", FAGFELT,
            index=FAGFELT.index(forslag),
        )

    if not jurynavn:
        st.warning("Skriv inn navnet ditt for å begynne å vurdere.")
        return

    kriterier_denne = KRITERIER_PER_FAGFELT[fagfelt_valgt]

    alle_rader = hent_vurderinger(sh)
    mine_vurderinger = {}
    if len(alle_rader) > 1:
        for rad in alle_rader[1:]:
            if len(rad) >= 6 and rad[1] == jurynavn:
                mine_vurderinger[(rad[0], rad[3])] = {"score": rad[4], "kommentar": rad[5]}

    sok = st.text_input("🔍 Søk etter butikk", "")
    navn_liste = sorted(butikkliste.keys())
    if sok:
        navn_liste = [n for n in navn_liste if sok.lower() in n.lower()]

    antall_fullfort = sum(1 for n in butikkliste if all((n, krit) in mine_vurderinger for krit in kriterier_denne))
    st.caption(f"**{fagfelt_valgt}** · {antall_fullfort} av {len(butikkliste)} butikker fullført av deg")
    st.progress(antall_fullfort / len(butikkliste) if butikkliste else 0)

    for navn in navn_liste:
        info = butikkliste[navn]
        alt_ferdig = all((navn, krit) in mine_vurderinger for krit in kriterier_denne)
        merke = "✅ " if alt_ferdig else ""
        with st.expander(f"{merke}**{navn}** — {info.get('klasse', '–')} · {info.get('bransje', '–')}"):
            if info.get("url"):
                st.markdown(f"🌐 [Besøk nettbutikk]({info['url']})")
            st.markdown("---")
            for krit in kriterier_denne:
                eksisterende = mine_vurderinger.get((navn, krit), {})
                st.markdown(f"**{krit}**")
                kc1, kc2 = st.columns([1, 2])
                with kc1:
                    score = st.select_slider(
                        "Score", [1, 2, 3, 4, 5],
                        value=int(eksisterende.get("score") or 3),
                        format_func=lambda x: "⭐" * x,
                        key=f"score_{navn}_{krit}",
                        label_visibility="collapsed",
                    )
                with kc2:
                    st.text_area(
                        "Kommentar", value=eksisterende.get("kommentar", ""),
                        key=f"kom_{navn}_{krit}", label_visibility="collapsed",
                        placeholder="Skriv en kommentar (valgfritt)…", height=68,
                    )
                st.markdown("")
            if st.button("💾 Lagre vurdering for denne butikken", key=f"lagre_{navn}"):
                vurderinger = [
                    (krit, st.session_state.get(f"score_{navn}_{krit}", 3), st.session_state.get(f"kom_{navn}_{krit}", ""))
                    for krit in kriterier_denne
                ]
                lagre_vurderinger_batch(sh, navn, jurynavn, fagfelt_valgt, vurderinger)
                st.success(f"Lagret vurdering for {navn}!")
                st.rerun()


# ─────────────────────────────────────────────
# Side: Finale
# ─────────────────────────────────────────────
def vis_finale(sh, butikkliste):
    st.header("🏆 Finale")
    st.caption("Rangering basert på snitt av alle registrerte ekspertvurderinger (Fase 2), gruppert per størrelsesklasse.")

    if not sh:
        st.warning("Google Sheets-tilkoblingen mangler – sjekk `secrets` i sidepanelet.")
        return
    if not butikkliste:
        st.info("💡 Last opp Fase 1-Excel-fila i sidepanelet til venstre for å se Finale-rangeringen (klasse hentes derfra).")
        return

    antall_per_klasse = st.number_input("Antall finalister som vises per klasse", min_value=1, max_value=20, value=3)

    alle_rader = hent_vurderinger(sh)
    snitt, detaljer = beregn_finale_data(alle_rader, butikkliste)

    if not snitt:
        st.info("Ingen vurderinger registrert ennå.")
        return

    for klasse in KLASSER:
        st.subheader(f"Klasse: {klasse}")
        butikker_i_klasse = [
            (navn, snitt[navn][0], snitt[navn][1])
            for navn in butikkliste
            if butikkliste[navn].get("klasse") == klasse and navn in snitt
        ]
        butikker_i_klasse.sort(key=lambda x: x[1], reverse=True)
        topp = butikker_i_klasse[:antall_per_klasse]

        if not topp:
            st.caption("Ingen vurderte butikker i denne klassen ennå.")
            continue

        for plass, (navn, snittscore, antall) in enumerate(topp, start=1):
            st.markdown(
                f'<div class="finale-rad"><span class="finale-plass">#{plass}</span>'
                f'<strong>{navn}</strong> &nbsp; <span class="snitt-badge">⭐ {snittscore:.2f} snitt</span>'
                f' &nbsp; <span style="color:#777;font-size:13px;">({antall} registrerte vurderinger)</span></div>',
                unsafe_allow_html=True,
            )
            with st.expander(f"Se alle vurderinger og kommentarer for {navn}"):
                info = butikkliste.get(navn, {})
                if info.get("url"):
                    st.markdown(f"🌐 [Besøk nettbutikk]({info['url']})")
                for fagfelt, kriterium, score, kommentar, jurymedlem in sorted(detaljer[navn]):
                    st.markdown(f"**{kriterium}** _({fagfelt} · {jurymedlem})_ — {'⭐' * int(score)}")
                    if kommentar:
                        st.caption(kommentar)


if side == "📖 Oversikt":
    vis_oversikt()
elif side == "✍️ Vurdering":
    vis_vurdering(sh, butikkliste)
elif side == "🏆 Finale":
    vis_finale(sh, butikkliste)

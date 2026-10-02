"""Domenekonstanter, Excel-import og ren beregningslogikk – ingen Google
Sheets-kall her (det ligger i sheets.py). Selve kriterietekstene og
poengberegningen er UENDRET fra tidligere versjoner."""

import re

import openpyxl
import streamlit as st

# ─────────────────────────────────────────────
# Kriterier – IDENTISK ordlyd som i Fase 1-arket (5 kategorier, 16 kriterier)
# ─────────────────────────────────────────────
KRITERIER = [
    ("Førsteinntrykk", "Opplevelse: Bildebruk, beskrivelser o.l."),
    ("Førsteinntrykk", "Navigasjon / UX"),
    ("Førsteinntrykk", "Søk: auto-korrektur, synonymer, utlisting, forslag"),
    ("Kundeservice & Tilgjengelighet", "Lett tilgjengelig info om levering, retur og kjøpsvilkår"),
    ("Kundeservice & Tilgjengelighet", "Google/Trustpilot/andre åpne løsninger for kundetilfredshet"),
    ("Kundeservice & Tilgjengelighet", "Tilgjengelighet, åpningstider, kanaler, responstid"),
    ("Kundeservice & Tilgjengelighet", "Bærekraft"),
    ("Kjøp/inspirasjon/personalisering", "Kassen: Levering, betaling"),
    ("Kjøp/inspirasjon/personalisering", "Inspirasjon"),
    ("Kjøp/inspirasjon/personalisering", "Mersalg/anbefalinger/personalisering"),
    ("Markedsføring/kundedialog", "Bruk av SoMe"),
    ("Markedsføring/kundedialog", "Betalt og organisk synlighet"),
    ("Markedsføring/kundedialog", "E-post"),
    ("Innovasjon", "Innovative løsninger"),
    ("Innovasjon", "Adopsjon av nye teknologier for styrke konkurransekraft"),
    ("Innovasjon", "Kommersielt håndverk"),
]
FAGFELT = list(dict.fromkeys(kat for kat, _ in KRITERIER))
KRITERIER_PER_FAGFELT = {kat: [krit for k, krit in KRITERIER if k == kat] for kat in FAGFELT}

# Generisk skalaforklaring – vises som hjelpetekst under hvert kriterium. Samme
# skala for alle kriterier; si ifra om dere ønsker kriteriespesifikke eksempler
# i tillegg, så legger vi det til som en overstyring per kriterium.
SKALA_FORKLARING = (
    "1 = svak – mangler i stor grad eller fungerer dårlig.  \n"
    "3 = middels – til stede, men med tydelig forbedringspotensial.  \n"
    "5 = utmerket – blant de beste i konkurransen på dette punktet."
)

KLASSE_DEFINISJON = [
    ("Liten", "Under 50 mill kr"),
    ("Medium", "50–250 mill kr"),
    ("Stor", "Over 250 mill kr"),
]
KLASSER = [k for k, _ in KLASSE_DEFINISJON]

# Forslag til fagfelt per jurymedlem – brukes KUN til å så Jury-fanen i Sheets
# første gang (se sheets.setup_from_excel). Etter det er Jury-fanen fasiten og
# redigeres manuelt der, denne lista leses ikke igjen.
JURY_FAGFELT_FORSLAG = {
    "Ole Johan": "Førsteinntrykk",
    "Stian": "Kundeservice & Tilgjengelighet",
    "Torkel": "Kjøp/inspirasjon/personalisering",
    "Marte": "Kjøp/inspirasjon/personalisering",
    "Vikki": "Markedsføring/kundedialog",
    "Nicholas": "Markedsføring/kundedialog",
    "Guro": "Innovasjon",
}

IKKE_VURDERT = "N/A"


# ─────────────────────────────────────────────
# Import av Fase 1-Excel-fila
# ─────────────────────────────────────────────
def les_fase1_liste(opplastet_fil):
    """Leser butikkliste fra Fase 1-arket: Butikk, Klasse, Bransje, URL, og
    eventuell «Snitt totalt»-kolonne (Fase 1-score, hvis jury har fylt den ut).
    URL-kolonnen kan enten være ren tekst eller en =HYPERLINK(...)-formel."""
    try:
        wb = openpyxl.load_workbook(opplastet_fil, data_only=False)
    except Exception as e:
        st.error(f"Kunne ikke lese Excel-filen: {e}", icon=":material/error:")
        return None

    ws = wb.worksheets[0]
    header_rad = None
    for r in range(1, min(ws.max_row, 10) + 1):
        if str(ws.cell(row=r, column=1).value or "").strip().lower() == "butikk":
            header_rad = r
            break
    if header_rad is None:
        st.error('Fant ikke en rad med "Butikk" i kolonne A i fila.', icon=":material/error:")
        return None

    kolonner = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_rad, column=c).value
        if v:
            kolonner[str(v).strip().lower()] = c

    kol_butikk = kolonner.get("butikk")
    kol_klasse = kolonner.get("klasse")
    kol_bransje = kolonner.get("bransje")
    kol_url = kolonner.get("url")
    kol_snitt = kolonner.get("snitt totalt")

    ws_verdier = None
    if kol_snitt:
        try:
            opplastet_fil.seek(0)
            ws_verdier = openpyxl.load_workbook(opplastet_fil, data_only=True).worksheets[0]
            opplastet_fil.seek(0)
        except Exception:
            ws_verdier = None

    butikker = {}
    for r in range(header_rad + 1, ws.max_row + 1):
        navn_celle = ws.cell(row=r, column=kol_butikk).value
        if not navn_celle or not str(navn_celle).strip():
            continue
        navn = str(navn_celle).strip()

        url = ""
        if kol_url:
            url_verdi = ws.cell(row=r, column=kol_url).value
            if url_verdi:
                treff = re.search(r'HYPERLINK\("([^"]+)"', str(url_verdi))
                url = treff.group(1) if treff else str(url_verdi).strip()

        fase1_snitt = None
        if ws_verdier is not None:
            snitt_verdi = ws_verdier.cell(row=r, column=kol_snitt).value
            if isinstance(snitt_verdi, (int, float)):
                fase1_snitt = float(snitt_verdi)

        butikker[navn] = {
            "klasse": str(ws.cell(row=r, column=kol_klasse).value or "").strip() if kol_klasse else "",
            "bransje": str(ws.cell(row=r, column=kol_bransje).value or "").strip() if kol_bransje else "",
            "url": url,
            "fase1_snitt": fase1_snitt,
        }
    if not butikker:
        st.error('Fant raden med "Butikk", men ingen butikker under den. Sjekk at fila har data under header-raden.', icon=":material/error:")
        return None
    return butikker


# ─────────────────────────────────────────────
# Beregning – fullføringsstatus og snitt, med støtte for "Kan ikke vurdere"
# ─────────────────────────────────────────────
def er_tall(score) -> bool:
    try:
        float(score)
        return True
    except (ValueError, TypeError):
        return False


def status_for_butikk(butikk_id, kriterier, vurderinger_oppslag):
    """vurderinger_oppslag: {(butikk_id, kriterium): {"score":.., "kommentar":..}}.
    Returnerer "Ferdig", "Påbegynt" eller "Ikke startet"."""
    antall_besvart = sum(1 for k in kriterier if (butikk_id, k) in vurderinger_oppslag)
    if antall_besvart == 0:
        return "Ikke startet"
    if antall_besvart == len(kriterier):
        return "Ferdig"
    return "Påbegynt"


def snitt_for_butikk(butikk_id, kriterier, vurderinger_oppslag):
    """Snitt av tallscorer for kriteriene (N/A/«Kan ikke vurdere» telles ikke med).
    Returnerer None hvis ingen tallscorer er gitt ennå."""
    tall = [
        float(vurderinger_oppslag[(butikk_id, k)]["score"])
        for k in kriterier
        if (butikk_id, k) in vurderinger_oppslag and er_tall(vurderinger_oppslag[(butikk_id, k)]["score"])
    ]
    return (sum(tall) / len(tall)) if tall else None
